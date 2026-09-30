from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document

from app.models.report import Report
from app.services import document_generator
from app.services.document_generator import DocumentGenerator

TITLE = "CONDIÇÕES DA VISTORIA INICIAL"
CLAUSES = (
    "• O imóvel está sendo entregue com as instalações elétricas e hidráulicas em perfeito "
    "funcionamento, apresentando-se em boas condições de higiene, limpeza e conservação, com "
    "todos os cômodos e paredes pintados, sendo que portas, portões e acessórios se encontram "
    "também em funcionamento correto, devendo o LOCATÁRIO mantê-lo desta forma, salvo as "
    "observações descritas acima;",
    "• No ato da entrega definitiva das chaves, o LOCATÁRIO restituirá o imóvel locado nas "
    "mesmas condições as quais o recebeu, ou seja, pintado ou as instalações elétricas, "
    "hidráulicas e acessórios deverão também estar em perfeitas condições de funcionamento, "
    "salvo as deteriorações decorrentes do uso normal e habitual do imóvel;",
    "• Após a assinatura, o LOCATÁRIO têm um prazo de 10 (dez) dias para apresentar qualquer "
    "irregularidade contida neste documento; passado este prazo, o LOCATÁRIO está ciente que "
    "não poderá efetuar reclamações posteriores.",
)


@pytest.mark.parametrize("report_type", ["Inicial", "Final"])
def test_conditions_only_for_exact_initial_type(tmp_path: Path, report_type: str) -> None:
    report = Report(report_type=report_type)
    before = deepcopy(report)
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(report, tmp_path, destination)

    document = Document(destination)
    paragraphs = [p.text for p in document.paragraphs]
    if report_type == "Inicial":
        start = paragraphs.index(TITLE)
        assert paragraphs[start + 1:start + 4] == list(CLAUSES)
        assert paragraphs[start + 4] == "6. TERMOS FINAIS"
        for paragraph in document.paragraphs[start + 1:start + 4]:
            assert all(run.bold for run in paragraph.runs if "LOCATÁRIO" in run.text)
    else:
        assert TITLE not in paragraphs
        assert all(clause not in paragraphs for clause in CLAUSES)
    assert report == before


def test_conditions_are_editable_in_template(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    template = Document(document_generator.TEMPLATE_PATH)
    clause = next(p for p in template.paragraphs if p.text == CLAUSES[0])
    source = Path(document_generator.__file__).read_text(encoding="utf-8")
    assert all(text not in source for text in CLAUSES)
    clause.text = "Texto de teste editado somente no template."
    custom_template = tmp_path / "template.docx"
    template.save(custom_template)
    monkeypatch.setattr(document_generator, "TEMPLATE_PATH", custom_template)
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(Report(report_type="Inicial"), tmp_path, destination)

    assert clause.text in [p.text for p in Document(destination).paragraphs]
