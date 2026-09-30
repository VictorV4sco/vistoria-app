from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from PIL import Image

from app.models.photo import Photo
from app.models.report import Report
from app.models.section import Section
from app.services.document_generator import DocumentGenerator
from app.services.image_service import ImageService


@pytest.mark.parametrize("counts", [(0,), (1,), (2,), (3,), (2, 2), (1, 0, 2)])
def test_photo_tables_content_order_optimization_and_proportions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, counts: tuple[int, ...],
) -> None:
    project = tmp_path / "project"
    originals = project / "imagens" / "originals"
    originals.mkdir(parents=True)
    report = Report()
    source_bytes = {}
    expected = []
    number = 0
    for index, count in enumerate(counts):
        section = Section(name=f"Ambiente {index}", order=10 - index)
        report.sections.append(section)
        for local_index in range(count):
            number += 1
            extension = "jpg" if number % 2 else "png"
            relative = Path("imagens/originals") / f"{number}.{extension}"
            size = (2400, 1200) if number % 2 else (1200, 2400)
            Image.new("RGB", size, (number * 40, 60, 100)).save(project / relative)
            source_bytes[relative] = (project / relative).read_bytes()
            caption = "Visão — janelas & <portão>" if number % 2 else ""
            section.photos.append(Photo(str(relative), caption=caption, order=10 - local_index))
            expected.append((relative, f"Foto {number:02d}" + (f" — {caption}" if caption else "")))
    before = deepcopy(report)
    calls = []
    optimize = ImageService.optimize_for_report

    def track_optimization(relative_path: Path, project_directory: Path) -> Path:
        result = optimize(relative_path, project_directory)
        calls.append((relative_path, project_directory, result))
        return result

    monkeypatch.setattr(ImageService, "optimize_for_report", track_optimization)
    monkeypatch.chdir(tmp_path)
    destination = tmp_path / "report.docx"
    DocumentGenerator.generate(report, project, destination)

    document = Document(destination)
    assert len(document.inline_shapes) == number + 1
    photo_tables = [table for table in document.tables if table._tbl.xpath(".//w:drawing")]
    assert len(photo_tables) == sum(count > 0 for count in counts)
    assert [p.text for p in document.paragraphs].count("Registro fotográfico") == sum(
        count > 0 for count in counts
    )
    assert [(relative, directory) for relative, directory, _ in calls] == [
        (relative, project) for relative, _ in expected
    ]
    photo_index = 0

    for table, count in zip(photo_tables, (c for c in counts if c), strict=True):
        assert len(table.columns) == 2
        assert len(table.rows) == (count + 1) // 2
        borders = table._tbl.xpath("./w:tblPr/w:tblBorders/*")
        assert borders and all(b.get(qn("w:val")) in ("nil", "none") for b in borders)
        for cell_index, cell in enumerate(cell for row in table.rows for cell in row.cells):
            if cell_index >= count:
                assert cell.text.strip() == ""
                assert not cell._tc.xpath(".//w:drawing")
                continue
            paragraphs = cell.paragraphs
            assert paragraphs[0]._p.xpath(".//w:drawing")
            assert paragraphs[1].text == expected[photo_index][1]
            blip = cell._tc.xpath(".//a:blip")[0]
            blob = document.part.related_parts[blip.get(qn("r:embed"))].blob
            optimized = project / calls[photo_index][2]
            assert blob == optimized.read_bytes()
            assert blob != source_bytes[expected[photo_index][0]]
            with Image.open(BytesIO(blob)) as image:
                assert max(image.size) <= 1600
                extent = cell._tc.xpath(".//wp:extent")[0]
                width, height = int(extent.get("cx")), int(extent.get("cy"))
                assert width / height == pytest.approx(image.width / image.height, rel=1e-5)
                assert width <= Cm(7)
                assert height <= Cm(6.5)
                assert paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
                assert paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER
            photo_index += 1
    with ZipFile(destination) as archive:
        assert archive.testzip() is None
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert len(media) == number + 1
    assert report == before
    assert all((project / path).read_bytes() == content for path, content in source_bytes.items())
    # Tables stay within their respective section, including sections without photos.
    blocks = []
    for child in document.element.body:
        if child.tag == qn("w:p"):
            text = "".join(child.itertext())
            if "Ambiente" in text:
                blocks.append("section")
        elif child.tag == qn("w:tbl") and child.xpath(".//w:drawing"):
            blocks.append("photos")
    assert blocks == [block for count in counts for block in
                      (["section", "photos"] if count else ["section"])]


def test_missing_photo_preserves_previous_document(tmp_path: Path) -> None:
    destination = tmp_path / "report.docx"
    Document().save(destination)
    previous = destination.read_bytes()
    report = Report(sections=[Section(name="Sala", photos=[Photo("missing.jpg")])])
    before = deepcopy(report)

    with pytest.raises(FileNotFoundError):
        DocumentGenerator.generate(report, tmp_path, destination)

    assert destination.read_bytes() == previous
    assert report == before
