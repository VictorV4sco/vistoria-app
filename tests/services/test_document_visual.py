from copy import deepcopy
from datetime import date
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from app.models.property import Property
from app.models.report import Report
from app.services.document_generator import DocumentGenerator


@pytest.mark.parametrize("pixels, expected", [
    ((2400, 1200), (6.0, 3.0)),
    ((1200, 2400), (4.0, 8.0)),
    ((1600, 1600), (6.0, 6.0)),
    ((96, 48), (2.54, 1.27)),
])
def test_photo_size_fits_box_without_upscaling(pixels: tuple, expected: tuple) -> None:
    from app.utils.image_dimensions import fit_photo_dimensions

    assert fit_photo_dimensions(*pixels) == pytest.approx(expected)


def test_cover_header_footer_and_styles(tmp_path: Path) -> None:
    report = Report(
        title="Vistoria fictícia", report_type="Inicial", inspection_date=date(2026, 9, 29),
        property=Property(property_type="Apartamento", address="Rua Fictícia", number="42"),
    )
    before = deepcopy(report)
    destination = tmp_path / "report.docx"
    DocumentGenerator.generate(report, tmp_path, destination)
    document = Document(destination)
    paragraphs = document.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text == "DADOS DA VISTORIA")
    cover = "\n".join(p.text for p in paragraphs[:start])
    for value in ("RELATÓRIO DE VISTORIA", "ALGER IMÓVEIS LTDA", "CRECI Nº 005103/0",
                  "Administração e venda de imóveis", "VISTORIA INICIAL", "Apartamento",
                  "Rua Fictícia", "42", "29/09/2026"):
        assert value in cover
    assert len(document.sections) == 2
    assert paragraphs[start - 1]._p.xpath('./w:pPr/w:sectPr')
    assert not paragraphs[start].paragraph_format.page_break_before
    assert any(p._p.xpath('.//w:drawing') for p in paragraphs[:start])
    section = document.sections[0]
    assert section.different_first_page_header_footer
    header = section.header
    header_text = "\n".join(p.text for p in header.paragraphs)
    assert "ALGER IMÓVEIS LTDA" in header_text
    assert "Relatório de Vistoria" in header_text
    assert "CRECI" not in header_text
    institutional_details = (
        "Rodovia Washington Luiz, nº 18.208, Loja C",
        "Santa Cruz da Serra, Duque de Caxias / RJ, Cep: 25.265-008",
        "CNPJ 04.960.485/0001-16",
        "(21) 97011-1382", "(21) 3658-6951", "(21) 99268-1605",
        "algerimoveis@gmail.com.br", "www.algerimoveis.com.br",
        "facebook.com/imoveisalger",
    )
    body = "\n".join(p.text for p in paragraphs[start:])
    for value in institutional_details:
        assert value in cover
        assert value not in header_text
        assert value not in body
    assert header._element.xpath('.//w:drawing')
    assert header._element.xpath('.//w:pBdr/w:bottom')
    fields = [node.text.strip() for node in section.footer._element.iter(qn('w:instrText'))]
    assert "PAGE" in fields and "NUMPAGES" in fields
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        assert document.styles[name].font.name == "Arial"
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        assert document.styles[name].paragraph_format.keep_with_next
    with ZipFile(destination) as archive:
        assert archive.testzip() is None
        logo = Path(__file__).resolve().parents[2] / "assets/logo_alger.png"
        assert logo.read_bytes() in [archive.read(name) for name in archive.namelist()
                                     if name.startswith("word/media/")]
    assert report == before


@pytest.mark.parametrize("report_type", ["Inicial", "Final"])
def test_centered_cover_and_closing_heading_pagination(tmp_path: Path, report_type: str) -> None:
    destination = tmp_path / "report.docx"
    DocumentGenerator.generate(Report(report_type=report_type), tmp_path, destination)
    document = Document(destination)
    start = next(i for i, p in enumerate(document.paragraphs) if p.text == "DADOS DA VISTORIA")
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in document.paragraphs[:start])
    assert len(document.sections) == 2
    cover, internal = document.sections
    # Center the entire cover 1.5 cm above its previous position.
    assert cover.top_margin == internal.top_margin
    assert cover.bottom_margin.cm - internal.bottom_margin.cm == pytest.approx(3.0, abs=0.01)
    assert cover._sectPr.xpath('./w:vAlign')[0].get(qn('w:val')) == 'center'
    assert internal._sectPr.xpath('./w:vAlign')[0].get(qn('w:val')) == 'top'
    assert internal._sectPr.xpath('./w:type')[0].get(qn('w:val')) == 'nextPage'
    assert not internal.different_first_page_header_footer
    assert internal.header._element.xpath('.//w:drawing')
    assert "ALGER IMÓVEIS LTDA" in "\n".join(p.text for p in internal.header.paragraphs)
    fields = [n.text.strip() for n in internal.footer._element.iter(qn('w:instrText'))]
    assert "PAGE" in fields and "NUMPAGES" in fields
    for p in document.paragraphs:
        if p.text in ("6. TERMOS FINAIS", "7. ASSINATURAS", "CONDIÇÕES DA VISTORIA INICIAL"):
            assert p.paragraph_format.keep_with_next
            assert p.paragraph_format.page_break_before is False
    terms_index = next(i for i, p in enumerate(document.paragraphs)
                       if p.text == "6. TERMOS FINAIS")
    assert document.paragraphs[terms_index + 1].paragraph_format.keep_together
