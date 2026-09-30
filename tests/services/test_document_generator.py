from copy import deepcopy
from datetime import date
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docxtpl import DocxTemplate

from app.models.party import Party
from app.models.property import Property
from app.models.report import Report
from app.models.section import Section
from app.services.document_generator import DocumentGenerator


@pytest.fixture
def report() -> Report:
    return Report(
        title="Vistoria — São Gonçalo & <Anexo>",
        report_type="Final",
        inspection_date=date(2026, 9, 4),
        inspector_name="João da Conceição",
        code="REF-ação-42",
        issue_date=date(2026, 10, 5),
        property=Property(
            property_type="Apartamento", description="Descrição — amplo & <arejado>",
            address="Rua do Ipê", number="123-B", complement="Fundos — 2º andar",
            neighborhood="Jardim América", city="Niterói", state="RJ", postal_code="24000-123",
        ),
        landlord=Party(
            name="José & <Sócios>", document="123.456.789-01", phone="(21) 99999-1111",
            email="jose@example.com", address="Rua São José, 10",
        ),
        tenant=Party(
            name="Lúcia Gonçalves", document="98.765.432/0001-10", phone="(21) 98888-2222",
            email="lucia@example.com", address="Avenida Açucena, 20",
        ),
    )


def test_generate_valid_docx_with_general_fields(
    tmp_path: Path, report: Report, monkeypatch: pytest.MonkeyPatch
) -> None:
    destination = tmp_path / "relatório.docx"
    before = deepcopy(report)
    monkeypatch.chdir(tmp_path)

    DocumentGenerator.generate(report, tmp_path / "project", destination)

    with ZipFile(destination) as archive:
        assert archive.testzip() is None
        assert "word/document.xml" in archive.namelist()
    document = Document(destination)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for value in (report.title, report.report_type, report.inspector_name, "04/09/2026"):
        assert value in text
    assert "{{" not in text
    assert report == before
    assert list(tmp_path.iterdir()) == [destination]


def test_all_fields_appear_in_the_correct_document_section(tmp_path: Path, report: Report) -> None:
    before = deepcopy(report)
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(report, tmp_path, destination)

    text = "\n".join(p.text for p in Document(destination).paragraphs)
    inspection, rest = text.split("DADOS DO IMÓVEL")
    property_text, parties = rest.split("PARTES ENVOLVIDAS")
    parties = parties.split("4. VISTORIA DOS AMBIENTES")[0]
    landlord_text, tenant_text = parties.split("LOCADOR")[1].split("LOCATÁRIO")
    assert "DADOS DA VISTORIA" in inspection
    for value in (report.title, report.report_type, report.code, report.inspector_name,
                  "04/09/2026", "05/10/2026"):
        assert value in inspection
    for value in vars(report.property).values():
        assert value in property_text
    for value in vars(report.landlord).values():
        assert value in landlord_text
        assert value not in tenant_text
    for value in vars(report.tenant).values():
        assert value in tenant_text
        assert value not in landlord_text
    assert "2026-09-04" not in text
    assert "2026-10-05" not in text
    assert report == before
    assert report.property == before.property
    assert report.landlord == before.landlord
    assert report.tenant == before.tenant


@pytest.mark.parametrize(
    ("owner", "field", "label"),
    [("report", "code", "Código / referência:"),
     ("property", "description", "Descrição resumida:"),
     ("property", "complement", "Complemento:"),
     ("property", "postal_code", "CEP:")]
    + [(party, field, label) for party in ("landlord", "tenant")
       for field, label in (("phone", "Telefone:"), ("email", "E-mail:"),
                            ("address", "Endereço:"))],
)
def test_empty_optional_field_omits_its_label(
    tmp_path: Path, report: Report, owner: str, field: str, label: str,
) -> None:
    setattr(report if owner == "report" else getattr(report, owner), field, "")
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(report, tmp_path, destination)

    text = "\n".join(p.text for p in Document(destination).paragraphs)
    if owner == "landlord":
        text = text.split("LOCADOR")[1].split("LOCATÁRIO")[0]
    elif owner == "tenant":
        text = text.split("LOCATÁRIO")[1]
    assert label not in text


def test_empty_report_has_no_empty_labels_or_none(tmp_path: Path) -> None:
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(Report(), tmp_path, destination)

    text = "\n".join(p.text for p in Document(destination).paragraphs)
    # Institutional contact labels on the cover are not optional report fields.
    assert ":" not in text.split("DADOS DA VISTORIA", 1)[1].replace(
        "Tipo da vistoria: Inicial", ""
    )
    assert "None" not in text
    assert "{{" not in text


def test_missing_inspection_date_does_not_render_none(tmp_path: Path, report: Report) -> None:
    report.inspection_date = None
    before = deepcopy(report)
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(report, tmp_path, destination)

    text = "\n".join(paragraph.text for paragraph in Document(destination).paragraphs)
    assert report.title in text
    assert "None" not in text
    assert "Data da vistoria:" not in text
    assert report == before


def test_success_replaces_existing_document(tmp_path: Path, report: Report) -> None:
    destination = tmp_path / "report.docx"
    Document().save(destination)

    DocumentGenerator.generate(report, tmp_path, destination)

    assert report.title in "\n".join(p.text for p in Document(destination).paragraphs)
    assert list(tmp_path.iterdir()) == [destination]


@pytest.mark.parametrize("count", [0, 1, 3])
def test_sections_follow_list_order_with_sequential_numbers(
    tmp_path: Path, report: Report, count: int,
) -> None:
    report.sections = [
        Section(name="ÁREA & <SERVIÇO>", description="Pintura íntegra & clara",
                notes="Atenção à ventilação", order=30),
        Section(name="COZINHA", description="Azulejos íntegros", order=10),
        Section(name="QUARTO — HÓSPEDES", notes="Maçaneta com folga", order=20),
    ][:count]
    before = deepcopy(report)
    original_sections = list(report.sections)
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(report, tmp_path, destination)

    with ZipFile(destination) as archive:
        assert archive.testzip() is None
    paragraphs = [p.text for p in Document(destination).paragraphs]
    start = paragraphs.index("4. VISTORIA DOS AMBIENTES")
    expected = []
    for index, section in enumerate(original_sections, start=1):
        expected.append(f"4.{index} {section.name}")
        if section.description:
            expected.extend(["Descrição", section.description])
        if section.notes:
            expected.extend(["Observações", section.notes])
    assert paragraphs[start + 1:paragraphs.index("6. TERMOS FINAIS")] == expected
    assert report == before
    assert all(current is original for current, original
               in zip(report.sections, original_sections, strict=True))


@pytest.mark.parametrize("description", ["", None, "Descrição — íntegra & <nova>"])
@pytest.mark.parametrize("notes", ["", None, "Observação — revisão necessária"])
def test_section_optional_text_and_labels(
    tmp_path: Path, description: str | None, notes: str | None,
) -> None:
    report = Report(
        report_type="Final", sections=[Section(name="Sala", description=description, notes=notes)]
    )
    before = deepcopy(report)
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(report, tmp_path, destination)

    paragraphs = [p.text for p in Document(destination).paragraphs]
    start = paragraphs.index("4.1 Sala")
    expected = []
    if description:
        expected.extend(["Descrição", description])
    if notes:
        expected.extend(["Observações", notes])
    assert paragraphs[start + 1:paragraphs.index("6. TERMOS FINAIS")] == expected
    assert "None" not in "\n".join(paragraphs)
    assert report == before


@pytest.mark.parametrize("existing", [False, True])
@pytest.mark.parametrize("stage", ["render", "save", "replace"])
def test_generation_failure_preserves_destination_and_cleans_temporary_files(
    tmp_path: Path, report: Report, monkeypatch: pytest.MonkeyPatch,
    existing: bool, stage: str,
) -> None:
    destination = tmp_path / "report.docx"
    if existing:
        Document().save(destination)
    before_files = {p.name: p.read_bytes() for p in tmp_path.iterdir()}
    before_report = deepcopy(report)

    def fail_render(self: DocxTemplate, *args: object, **kwargs: object) -> None:
        raise ValueError("render failed")

    def fail_save(self: DocxTemplate, path: Path, **kwargs: object) -> None:
        path = Path(path)
        assert path != destination
        assert path.is_relative_to(tmp_path)
        path.write_bytes(b"partial docx")
        raise OSError("save failed")

    def fail_replace(source: Path, target: Path) -> None:
        assert target == destination
        assert report.title in "\n".join(p.text for p in Document(source).paragraphs)
        raise PermissionError("replace failed")

    if stage == "render":
        monkeypatch.setattr(DocxTemplate, "render", fail_render)
    elif stage == "save":
        monkeypatch.setattr(DocxTemplate, "save", fail_save)
    else:
        monkeypatch.setattr(Path, "replace", fail_replace)

    with pytest.raises(ValueError if stage == "render" else OSError, match=f"{stage} failed"):
        DocumentGenerator.generate(report, tmp_path, destination)

    assert {p.name: p.read_bytes() for p in tmp_path.iterdir()} == before_files
    assert report == before_report
