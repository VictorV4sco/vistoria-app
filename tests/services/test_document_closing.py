from copy import deepcopy
from datetime import date
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from app.models.complementary_information import ComplementaryInformation
from app.models.party import Party
from app.models.report import Report
from app.services import document_generator
from app.services.document_generator import DocumentGenerator


@pytest.mark.parametrize("month, name", enumerate([
    "janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
    "setembro", "outubro", "novembro", "dezembro",
], start=1))
def test_portuguese_date_format(month: int, name: str) -> None:
    from app.utils.dates import format_date_pt_br

    assert format_date_pt_br(date(2026, month, 29 if month != 2 else 28)) == (
        f"{29 if month != 2 else 28} de {name} de 2026"
    )
    assert format_date_pt_br(None) == ""


@pytest.mark.parametrize("filled", [True, False, "partial"])
def test_complementary_information(tmp_path: Path, filled: bool | str) -> None:
    info = ComplementaryInformation(
        delivered_keys="3 chaves", energy_meter="Medidor — ação", consumer_unit="UC-123",
        general_notes="Revisão & <atenção>", issue_location="São Gonçalo/RJ",
    ) if filled else ComplementaryInformation()
    if filled == "partial":
        info.energy_meter = ""
        info.consumer_unit = ""
    report = Report(complementary_information=info)
    before = deepcopy(report)
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(report, tmp_path, destination)

    with ZipFile(destination) as archive:
        assert archive.testzip() is None
    text = "\n".join(p.text for p in Document(destination).paragraphs)
    if filled:
        text = text.split("5. INFORMAÇÕES COMPLEMENTARES")[1].split("6. TERMOS FINAIS")[0]
    else:
        assert "5. INFORMAÇÕES COMPLEMENTARES" not in text
    for field, label in [
        ("delivered_keys", "Chaves entregues"), ("energy_meter", "Medidor de energia"),
        ("consumer_unit", "Unidade consumidora"), ("general_notes", "Observações gerais"),
        ("issue_location", "Local de emissão"),
    ]:
        value = getattr(info, field)
        if value:
            assert f"{label}: {value}" in text
        else:
            assert f"{label}:" not in text
    assert report == before
    assert report.complementary_information is info


@pytest.mark.parametrize("location", ["Rio de Janeiro/RJ", ""])
@pytest.mark.parametrize("issue_date", [date(2026, 9, 29), None])
def test_closing_date_and_signatures(
    tmp_path: Path, location: str, issue_date: date | None,
) -> None:
    report = Report(
        issue_date=issue_date, inspector_name="Inspetor exclusivo",
        landlord=Party(name="José & <Sócios>"), tenant=Party(name="Lúcia Gonçalves"),
        complementary_information=ComplementaryInformation(issue_location=location),
    )
    before = deepcopy(report)
    destination = tmp_path / "report.docx"
    DocumentGenerator.generate(report, tmp_path, destination)
    paragraphs = [p.text for p in Document(destination).paragraphs]
    closing = paragraphs[paragraphs.index("6. TERMOS FINAIS") + 1:]
    if issue_date:
        expected = (location + ", " if location else "") + "29 de setembro de 2026"
        assert expected in closing
    else:
        assert not any("de 2026" in p for p in closing)
    document = Document(destination)
    assert "7. ASSINATURAS" in closing
    assert len(document.tables) == 1
    signatures = document.tables[0]
    assert len(signatures.rows) == 1 and len(signatures.columns) == 2
    for cell, title, name in zip(signatures.rows[0].cells,
                                 ("LOCADOR", "LOCATÁRIO"),
                                 (report.landlord.name, report.tenant.name), strict=True):
        assert [p.text for p in cell.paragraphs] == [
            title, "____________________________________", name,
        ]
        assert cell.paragraphs[1].paragraph_format.space_before.pt >= 24
        assert report.inspector_name not in cell.text
    assert "".join(signatures._tbl.getprevious().xpath(".//w:t/text()")) == "7. ASSINATURAS"
    assert "None" not in "\n".join(closing)
    assert report == before


def test_terms_are_editable_in_template(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    template = Document(document_generator.TEMPLATE_PATH)
    paragraphs = template.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text == "6. TERMOS FINAIS")
    terms = paragraphs[start + 1]
    assert len(terms.text) > 50
    assert "{{" not in terms.text and "{%" not in terms.text
    assert terms.text not in Path(document_generator.__file__).read_text(encoding="utf-8")
    replacement = "Texto de termos editado no template — revisão & conservação."
    terms.text = replacement
    custom_template = tmp_path / "template.docx"
    template.save(custom_template)
    monkeypatch.setattr(document_generator, "TEMPLATE_PATH", custom_template)
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(Report(), tmp_path, destination)

    assert replacement in [p.text for p in Document(destination).paragraphs]


@pytest.mark.parametrize("report_type", ["Inicial", "Final"])
@pytest.mark.parametrize("consumer_unit", ["UC-123", ""])
def test_report_type_controls_cover_and_energy_meter(
    tmp_path: Path, report_type: str, consumer_unit: str,
) -> None:
    report = Report(
        report_type=report_type,
        complementary_information=ComplementaryInformation(
            energy_meter="MEDIDOR-TESTE-42", consumer_unit=consumer_unit,
        ),
    )
    before = deepcopy(report)
    destination = tmp_path / "report.docx"

    DocumentGenerator.generate(report, tmp_path, destination)

    with ZipFile(destination) as archive:
        assert archive.testzip() is None
    text = "\n".join(p.text for p in Document(destination).paragraphs)
    cover = text.split("DADOS DA VISTORIA")[0]
    assert f"VISTORIA {report_type.upper()}" in cover
    if report_type == "Inicial":
        assert "CONDIÇÕES DA VISTORIA INICIAL" in text
        assert "Medidor de energia: MEDIDOR-TESTE-42" in text
    else:
        assert "CONDIÇÕES DA VISTORIA INICIAL" not in text
        assert "Medidor de energia" not in text
        assert "MEDIDOR-TESTE-42" not in text
        if not consumer_unit:
            assert "5. INFORMAÇÕES COMPLEMENTARES" not in text
    if consumer_unit:
        assert "Unidade consumidora: UC-123" in text
    assert report == before
    assert report.complementary_information.energy_meter == "MEDIDOR-TESTE-42"
